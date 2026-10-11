"""Rebuild editable Moniq slides; presentation tools live outside app dependencies.

Run with python-pptx and Pillow installed. Export with export_presentation.ps1.
"""
from pathlib import Path
import json
import re
from zipfile import ZipFile, ZIP_DEFLATED
import xml.etree.ElementTree as ET

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]
ASSETS = OUT / "assets"
CREAM, INK, GREEN, MUTED = "F6F7F3", "253C31", "176C54", "596B60"
LINE, SAGE, WHITE, DARK = "DCE5DC", "EAF1E7", "FFFFFF", "143E32"
FONT = "Segoe UI"
W, H = 13.333333, 7.5

SPEECH = [
    (10, "Meet Moniq: Spending, Made Smarter. It helps college students track spending, manage shared expenses, and understand their financial future."),
    (18, "Student money lives in fragments: bank alerts, spreadsheets, and separate apps. A paycheck lands, dinner gets split, and rent is coming. The challenge is connecting those events into a clear picture of what you can spend and what you need to save."),
    (18, "Moniq brings that picture together. Import transactions or connect Gmail with permission. Review the extracted activity, let machine learning suggest categories, and see confirmed records update your account estimates, budgets, and forecasts. Circles adds your personal share of group spending."),
    (30, "Here is the working flow, using synthetic data. Start with ten dollars. Moniq detects a four-hundred-dollar paycheck alert. After confirmation, the estimate becomes four hundred ten. A fifteen-dollar restaurant alert is extracted, categorized, and matched to checking. Confirm it, and the estimate becomes three hundred ninety-five. These are actual saved records and calculations, not numbers placed on a mockup. Gmail-derived balances remain estimates, not verified bank balances."),
    (26, "SmartApprove is optional. You control trusted sender rules; authentication, account, date, confidence, and amount checks decide whether an alert can skip review. Uncertain messages stay manual. Circles handles everyday roommate bookkeeping: pay forty dollars for dinner, keep the full charge in your cash ledger, and count twenty dollars as your share. A confirmed twenty-dollar repayment clears the debt without becoming paycheck income. Moniq never sends payments."),
    (25, "The system uses React and TypeScript over a Python FastAPI backend and SQLite. A local TF-IDF and logistic-regression model suggests categories. Deterministic code handles integer-cent balances, splits, and forecasts. An independent evaluation reproduced ninety-four point seven percent raw accuracy on held-out synthetic examples. That is evidence of the ML pipeline, not measured real-bank accuracy. Low-confidence descriptions still need review."),
    (18, "The prototype passed one hundred nine backend tests and nineteen Chromium browser tests, including production pages and isolated financial workflows. It includes responsive budgets, goals, subscriptions, and record-grounded answers. Wider real-email coverage, live bank verification, and production identity remain future work. Sample deals are fictional."),
    (15, "The goal is practical: less bookkeeping, clearer decisions, and easier roommate expenses. Codex assisted development; the work-history and dependency disclosures are in the repository. Moniq: your money, your goals, your financial future. Spending, Made Smarter."),
]


def color(value):
    return RGBColor.from_string(value)


def box(s, x, y, w, h, fill=WHITE, border=None, rounded=True):
    shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                               Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if border is None else None
    if border:
        shape.line.color.rgb = color(border)
        shape.line.width = Pt(0.8)
    if rounded:
        shape.adjustments[0] = 0.10
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    effect = shape._element.find('.//{http://schemas.openxmlformats.org/drawingml/2006/main}effectRef')
    if effect is not None:
        effect.set('idx', '0')
    return shape


def text(s, words, x, y, w, h, size=22, ink=INK, bold=False, align=None, link=None):
    shape = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for index, line in enumerate(words.split("\n")):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.12
        if align is not None:
            p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color(ink)
        if link:
            run.hyperlink.address = link
    return shape


def photo(s, filename, x, y, w, h):
    path = ASSETS / filename
    iw, ih = Image.open(path).size
    ratio = min(w / iw, h / ih)
    pw, ph = iw * ratio, ih * ratio
    box(s, x - .04, y - .04, w + .08, h + .08, WHITE, LINE)
    image = s.shapes.add_picture(str(path), Inches(x + (w - pw) / 2),
                                Inches(y + (h - ph) / 2), Inches(pw), Inches(ph))
    image._element.nvPicPr.cNvPr.set("descr", "Actual Moniq interface; isolated synthetic data. " + filename)


def slide(prs, number, chapter, title=None, dark=False):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color(DARK if dark else CREAM)
    c = "D1E4D7" if dark else MUTED
    text(s, chapter.upper(), .6, .35, 10, .28, 12, c, True)
    text(s, "moniq.", 11.25, .27, 1.47, .45, 22, WHITE if dark else GREEN, True, PP_ALIGN.RIGHT)
    box(s, .6, 7.02, 12.13, .012, "48705F" if dark else LINE, rounded=False)
    text(s, "Hacklanta 2026  /  Fintech", .6, 7.13, 5.9, .2, 10, c)
    text(s, f"{number:02d}", 12.05, 7.10, .67, .26, 12, c, align=PP_ALIGN.RIGHT)
    if title:
        text(s, title, .6, .95, 12.05, 1.24, 33, WHITE if dark else INK, True)
    return s


def notes(s, index, evidence):
    sec, words = SPEECH[index]
    s.notes_slide.notes_text_frame.text = f"SLIDE {index + 1} | {sec} seconds\n\n{words}\n\nEvidence / presenter reminders:\n{evidence}"


def prepare_assets():
    # Crops only: preserve the original screenshot/video content, no invented UI.
    crops = {
        "dashboard": (ROOT / "docs/screenshots/12-final-dashboard.png", (265, 110, 1820, 925)),
        "review": (ASSETS / "review-full.png", (405, 668, 1758, 909)),
        "balance": (ASSETS / "balance-full.png", (385, 306, 1221, 540)),
        "smartapprove": (ASSETS / "smartapprove-full.png", (405, 487, 1325, 775)),
        "circles": (ASSETS / "circles-full.png", (758, 428, 1775, 701)),
    }
    for name, (source, crop) in crops.items():
        Image.open(source).crop(crop).save(ASSETS / f"{name}.png")


def build_main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Moniq | Spending, Made Smarter. | Hacklanta 2026"
    prs.core_properties.author = "Moniq team — presenter to be confirmed"
    prs.core_properties.subject = "Fintech judging presentation with synthetic demonstration data"

    s = slide(prs, 1, "Spending, Made Smarter.")
    text(s, "moniq.", .6, 1.35, 4.35, 1.4, 82, GREEN, True)
    text(s, "Spending,\nMade Smarter.", .65, 2.93, 4.2, 1.25, 32, INK, True)
    text(s, "AI-powered personal finance\nfor college students.", .65, 4.52, 4.3, .95, 22, MUTED)
    photo(s, "dashboard.png", 5.28, 1.55, 7.45, 4.1)
    text(s, "ACTUAL MONIQ DASHBOARD  ·  SYNTHETIC DATA", 5.28, 5.90, 7.45, .3, 12, MUTED, True)
    text(s, "Presenter: [edit name / team]", .65, 6.40, 7.4, .35, 16, MUTED)
    notes(s, 0, "docs/screenshots/12-final-dashboard.png. Alex is a synthetic demo profile, not the presenter. Replace presenter placeholder. No customer/user-count claims.")

    s = slide(prs, 2, "The problem", "College finances are\nunnecessarily complicated.")
    for x, n, title, detail, label in [
        (.6, "01", "Money is scattered.", "Bank alerts, spreadsheets,\nand separate apps.", "FIND THE RECORDS"),
        (4.77, "02", "Tracking takes effort.", "New income and purchases\nare easy to miss.", "KEEP UP WITH ACTIVITY"),
        (8.94, "03", "Planning adds stress.", "Shared bills, rent,\nand savings goals.", "KNOW WHAT COMES NEXT"),
    ]:
        box(s, x, 2.68, 3.79, 3.47, WHITE, LINE)
        text(s, n, x + .25, 2.94, 1.2, .7, 38, GREEN, True)
        text(s, title, x + .25, 3.86, 3.32, .92, 22, INK, True)
        text(s, detail, x + .25, 4.83, 3.32, .9, 19, MUTED)
        text(s, label, x + .25, 5.80, 3.32, .25, 10, MUTED, True)
    notes(s, 1, "Problem framing is a use case, not a survey or market-size claim. No quantified impact or testimonials are asserted.")

    s = slide(prs, 3, "Our solution", "One intelligent home for your money.")
    for y, n, head, body in [(2.30, "1", "Bring activity together", "Gmail with consent · CSV imports"),
                             (3.55, "2", "Review and organize", "ML categories · Confirmed ledger"),
                             (4.80, "3", "Plan your next decision", "Budgets · Goals · Forecasts · Circles")]:
        box(s, .62, y, .50, .50, SAGE)
        text(s, n, .62, y + .05, .50, .35, 20, GREEN, True, PP_ALIGN.CENTER)
        text(s, head, 1.36, y, 4.4, .5, 23, INK, True)
        text(s, body, 1.36, y + .58, 4.5, .72, 17, MUTED)
    photo(s, "dashboard.png", 6.22, 2.18, 6.5, 3.8)
    text(s, "Confirmed activity refreshes the dashboard.", 6.22, 6.14, 6.5, .35, 17, GREEN, True)
    notes(s, 2, "README; frontend/src/pages; backend/app/finance.py, goals.py, gmail_routes.py, circles.py. 'Organize income' includes rules/event detection; ML is category classification, not deposit verification.")

    s = slide(prs, 4, "The wow moment", "From an email alert to an updated balance.")
    text(s, "SYNTHETIC DEMONSTRATION  ·  ACTUAL PARSER + DATABASE UPDATES", .62, 1.89, 12, .35, 13, MUTED, True)
    for x, value, caption, fill in [(.6, "$10", "Starting estimate", WHITE), (5.04, "$410", "Paycheck confirmed", WHITE), (9.48, "$395", "Purchase confirmed", GREEN)]:
        box(s, x, 2.43, 3.25, 1.43, fill, LINE if fill == WHITE else None)
        text(s, value, x + .25, 2.58, 2.75, .75, 42, WHITE if fill == GREEN else GREEN, True)
        text(s, caption, x + .25, 3.40, 2.8, .3, 15, WHITE if fill == GREEN else MUTED)
    text(s, "+ $400", 3.90, 2.71, 1.10, .34, 16, GREEN, True, PP_ALIGN.CENTER)
    text(s, "→", 3.90, 3.13, 1.10, .44, 22, GREEN, True, PP_ALIGN.CENTER)
    text(s, "− $15", 8.33, 2.71, 1.10, .34, 16, GREEN, True, PP_ALIGN.CENTER)
    text(s, "→", 8.33, 3.13, 1.10, .44, 22, GREEN, True, PP_ALIGN.CENTER)
    photo(s, "review.png", .65, 4.18, 7.70, 1.90)
    photo(s, "balance.png", 8.65, 4.42, 4.03, 1.36)
    text(s, "Detect → Extract → Categorize → Assign account → Confirm → Recalculate", .65, 6.30, 12.05, .35, 17, GREEN, True)
    text(s, "Gmail-derived balances are estimates, not verified bank balances.", .65, 6.71, 12.05, .25, 13, MUTED)
    notes(s, 3, "Video Moniq_Hacklanta_Final.mp4: purchase review 00:29 and confirmed dashboard 00:41. Synthetic account ending 1234. Financial email detection/extraction precedes ML; confirmation gate precedes ledger update. Synthetic scenario doesn't read actual Gmail.")

    s = slide(prs, 5, "SmartApprove + Moniq Circles", "Automation meets everyday student life.")
    box(s, .6, 2.12, 5.92, 4.22, WHITE, LINE)
    box(s, 6.81, 2.12, 5.92, 4.22, WHITE, LINE)
    text(s, "SmartApprove", .85, 2.35, 5.42, .45, 25, GREEN, True)
    text(s, "Your sender rules. Your choice.\nAuthentication + amount checks.\nUncertain alerts stay in review.", .85, 3.01, 5.42, 1.12, 19)
    photo(s, "smartapprove.png", .89, 4.40, 5.33, 1.26)
    text(s, "OPT-IN  ·  OFF BY DEFAULT", .85, 5.90, 5.42, .27, 12, MUTED, True)
    text(s, "Moniq Circles", 7.06, 2.35, 5.42, .45, 25, GREEN, True)
    text(s, "Split expenses; keep the full charge.\nTrack personal shares and repayments.\nReimbursements are not paycheck income.", 7.06, 3.01, 5.42, 1.12, 18)
    photo(s, "circles.png", 7.1, 4.25, 5.33, 1.52)
    text(s, "$40 dinner = $20 your share + $20 owed", 7.06, 5.94, 5.42, .27, 14, GREEN, True)
    text(s, "Synthetic example. Circles records debts and repayments; Moniq sends no money.", .65, 6.58, 12.0, .31, 14, MUTED)
    notes(s, 4, "backend/app/smartapprove.py checks aligned Google-reported authentication, posted status/date, account, amount/daily caps, confidence, dedup. Test-only sender authentication allowed solely in explicit controlled fixtures. Circles integer-cent allocation/settlement, not live collaboration/payment rails.")

    s = slide(prs, 6, "The technology", "Real machine learning. Real financial logic.")
    labels = ["Gmail API / User imports", "Secure FastAPI backend", "Financial event extraction", "Scikit-learn classification", "Validation + account ledger", "React dashboard, forecasts + Circles"]
    for i, label in enumerate(labels):
        y = 2.05 + i * .72
        box(s, .62, y, 6.02, .52, GREEN if i == 3 else WHITE, LINE if i != 3 else None)
        text(s, label, .83, y + .065, 5.60, .35, 18, WHITE if i == 3 else INK, True)
        if i < 5:
            arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.58), Inches(y + .56), Inches(.12), Inches(.12))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = color(GREEN)
            arrow.line.fill.background()
    text(s, "ML suggests categories.", 7.15, 2.08, 5.4, .5, 23, GREEN, True)
    text(s, "TF-IDF + logistic regression\n94.70% raw synthetic accuracy", 7.15, 2.76, 5.4, .95, 21)
    text(s, "250 / 264 held-out examples.\nReal-bank accuracy is unmeasured.", 7.15, 3.79, 5.4, .75, 16, MUTED)
    text(s, "Code calculates the money.", 7.15, 4.83, 5.4, .5, 23, GREEN, True)
    text(s, "Integer cents · Deterministic forecasts", 7.15, 5.43, 5.4, .45, 18)
    text(s, "React / TypeScript · Python / FastAPI · SQLite\nGoogle OAuth + Gmail API · pytest / Playwright", 7.15, 6.13, 5.4, .68, 14, MUTED)
    notes(s, 5, "backend/app/ml.py; scripts/reproduce_ml.py --check; data/ml_independent_evaluation.json. 792 train, 264 test, disjoint merchant families, correlated synthetic formatting, uncalibrated confidence. Threshold .45 gives 58.71% automatic coverage. No generative LLM is used for financial arithmetic or assistant replies. Backend security is local prototype scope, not certification.")

    s = slide(prs, 7, "Working prototype + validation", "Built, tested, and demonstrated.")
    for x, value, label, detail in [(.6, "109", "Backend tests passed", "pytest"), (4.77, "19", "Browser tests passed", "16 local + 3 production · Chromium"), (8.94, "94.70%", "Raw synthetic ML accuracy", "250 / 264 · Held-out merchant families")]:
        box(s, x, 2.08, 3.79, 1.65, WHITE, LINE)
        text(s, value, x + .24, 2.20, 3.30, .70, 41, GREEN, True)
        text(s, label, x + .24, 3.00, 3.30, .31, 16, INK, True)
        text(s, detail, x + .24, 3.39, 3.30, .24, 10, MUTED)
    text(s, "Verified prototype", .65, 4.08, 5.70, .5, 24, GREEN, True)
    text(s, "Responsive dashboard and budgets\nGoals, forecasts, subscriptions, assistant\nSynthetic email → confirmation → ledger\nSharing, repayment and session isolation", .65, 4.82, 5.7, 1.55, 18)
    text(s, "Clear boundaries", 7.06, 4.08, 5.5, .5, 24, GREEN, True)
    text(s, "Estimates; no bank verification\nGmail requires permission\nReal-email coverage needs broader testing\nNo payments; sample deals are fictional", 7.06, 4.82, 5.5, 1.55, 18)
    text(s, "Hackathon prototype  ·  Not a production banking product", .65, 6.60, 12.0, .3, 14, MUTED)
    notes(s, 6, "docs/QA_REPORT.md figures independently rerun for this presentation: 109 backend, 16 local Chromium, 3 public production Chromium; latest results in PRESENTATION_VERIFICATION.md. Counts reflect local validation. Final hosted GitHub CI remains pending approval to push.")

    s = slide(prs, 8, "The vision", dark=True)
    text(s, "Make financial confidence\naccessible to every student.", .65, 1.19, 12.0, 1.55, 39, WHITE, True)
    text(s, "Less bookkeeping. Clearer decisions.\nSmarter savings. Easier roommate expenses.", .68, 3.01, 11.95, 1.12, 26, "D1E4D7")
    text(s, "moniq.", .65, 4.52, 4.7, .95, 55, WHITE, True)
    text(s, "Spending, Made Smarter.", .68, 5.68, 9.8, .48, 25, WHITE)
    url = "https://github.com/ReddyR1105/Moniq"
    text(s, url, .68, 6.30, 12.0, .40, 18, "D1E4D7", link=url)
    text(s, "AI-assisted build · Work history disclosed in repository", .68, 6.77, 12.0, .22, 11, "D1E4D7")
    notes(s, 7, "Aspirations, not proven impact/user adoption. Repository URL points to published baseline; final source/presentation require approved push. Codex assisted code, tests, docs, recordings and presentation. Team must complete WORK_DISCLOSURE.md work-window/contribution fields. Do not assert all work happened during Hacklanta until confirmed.")
    path = OUT / "Moniq_Hacklanta_2026_Pitch.pptx"
    prs.save(path)
    # PowerPoint uses the theme's hyperlink color even with an explicit run color.
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    for name, data in entries.items():
        if name.startswith("ppt/theme/") and name.endswith(".xml"):
            tree = ET.fromstring(data)
            for key in ("hlink", "folHlink"):
                element = tree.find(".//a:clrScheme/a:" + key, ns)
                for child in list(element):
                    element.remove(child)
                ET.SubElement(element, "{" + ns["a"] + "}srgbClr", val="D1E4D7")
            entries[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return prs, path


def build_backup():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    prs.core_properties.title = "Moniq | Technical backup slides"
    s = slide(prs, 1, "Technical backup · outside timed pitch", "ML evidence, with its limits.")
    text(s, "792 train / 264 test", .65, 2.30, 6.0, .6, 31, GREEN, True)
    text(s, "11 categories · Disjoint merchant families\nWord uni/bigrams → logistic regression\n250 correct · Raw macro F1 0.9432", .65, 3.2, 6.0, 1.65, 23)
    text(s, "Confidence policy", 7.1, 2.35, 5.6, .6, 28, GREEN, True)
    text(s, "Below 45% → miscellaneous + review\n58.71% automatic coverage on this set\nConfidence is uncalibrated\nReal-bank performance not measured", 7.1, 3.2, 5.6, 2.25, 20)
    text(s, "Synthetic English examples have correlated formatting. No production accuracy claim.", .65, 6.27, 12.0, .5, 17, MUTED)
    s.notes_slide.notes_text_frame.text = "Evidence: scripts/reproduce_ml.py, data/ml_independent_evaluation.json, backend/app/ml.py. Raw accuracy precedes fallback. Accepted predictions 155/155 correct on this set only; miscellaneous fallback overall 63.64%. Do not hide review coverage or generalize synthetic results."
    s = slide(prs, 2, "Technical backup · outside timed pitch", "Permission, validation, exact-cent accounting.")
    for x, title, body in [(.65, "Gmail + privacy", "OAuth code + PKCE + one-use state\nGmail read-only scope\nEncrypted backend tokens\nPrivate / public / scenario isolation"), (7.0, "Ledger + recovery", "Posted records, integer cents\nDuplicate checks + idempotent requests\nAudit history + compensating reversals\nCircles preserves cash charge")]:
        text(s, title, x, 2.28, 5.5, .6, 27, GREEN, True)
        text(s, body, x, 3.19, 5.65, 2.5, 21)
    text(s, "Local owner prototype; SQLite is unencrypted. Cloud identity and security review remain future work.", .65, 6.28, 12.0, .56, 16, MUTED)
    s.notes_slide.notes_text_frame.text = "Evidence: gmail_service.py, gmail_routes.py, smartapprove.py, transactions.py, circles.py, account_tools.py. Provider authentication is not cryptographic financial proof. No bank payments/live collaboration. Work-window eligibility remains manual."
    path = OUT / "Moniq_Technical_Backup.pptx"
    prs.save(path)
    return path


def write_scripts():
    total = sum(seconds for seconds, _ in SPEECH)
    words = sum(len(re.findall(r"\b[\w'-]+\b", speech)) for _, speech in SPEECH)
    script = ["# Moniq judging pitch", "", "Eight slides; target **2:40** including brief slide transitions. Speaker notes in the PowerPoint contain the same script.", "", "Rehearse at about 150 words per minute. Do not read the fine print aloud. The slide-four workflow is illustrated with real synthetic recording frames; a live click-through is optional and replaces that section rather than extending the pitch.", "", "Presenter/team: **[edit name / team]**. Never use the synthetic Alex profile as the team roster.", ""]
    cursor = 0
    for i, (seconds, speech) in enumerate(SPEECH, 1):
        script.extend([f"## Slide {i} — {seconds} seconds ({cursor // 60}:{cursor % 60:02d}–{(cursor + seconds) // 60}:{(cursor + seconds) % 60:02d})", "", speech, ""])
        cursor += seconds
    script.extend(["## Rehearsal and demo", "", "- Open slide show at 1920×1080; check the projector before judging.", "- Keep `../demo/Moniq_Hacklanta_Final.mp4` (2:21) available as the separate required demo video; do not play the whole video inside this 2:40 pitch.", "- For live demonstration, use a dedicated isolated synthetic workspace with Gmail disabled. Follow `../DEMO_SCRIPT.md`; never reset the real owner workspace or open personal Gmail.", "- Confirm the official judging slot with organizers; no official pitch duration was established from the published pages.", "- Complete `MANUAL_SUBMISSION_CHECKLIST.md` and verified work/contribution disclosures before submission.", ""])
    (OUT / "PRESENTATION_SCRIPT.md").write_text("\n".join(script), encoding="utf-8")
    (OUT / "speaking-timing.json").write_text(json.dumps({"spoken_words": words, "allocated_seconds": total, "target_seconds_with_transitions": 160, "estimated_seconds_at_150_wpm": round(words / 150 * 60, 1), "estimated_seconds_at_130_wpm": round(words / 130 * 60, 1)}, indent=2), encoding="utf-8")
    quick = """# Moniq — 60-second pitch

**0–14 seconds:** Student money is scattered across bank alerts, spreadsheets, and roommate expenses. Moniq brings it into one clear picture: Spending, Made Smarter. Budgets, upcoming bills, goals, and forecasts help you plan ahead.

**14–31 seconds:** Here is our synthetic demo: start with ten dollars, confirm a four-hundred-dollar paycheck, then confirm a fifteen-dollar restaurant purchase. The estimate becomes three hundred ninety-five. Those are saved records and actual ledger updates, not a mockup.

**31–52 seconds:** Local machine learning suggests categories; deterministic code calculates money. SmartApprove offers optional checked automation. Circles splits a forty-dollar dinner into your twenty-dollar share and twenty dollars owed, without calling repayments paycheck income. Independent evaluation reproduced ninety-four point seven percent raw classification accuracy on synthetic examples; broader real-bank accuracy has not been measured.

**52–60 seconds:** We passed one hundred nine backend and nineteen browser tests. Balances are estimates; Moniq sends no payments. Your money, your goals, your financial future.

## Demo cue

Show slide 4 or the already prepared isolated synthetic dashboard. If time permits, replace the middle spoken example with actual confirmation clicks following `../DEMO_SCRIPT.md`. Do not attempt OAuth authorization during the one-minute pitch. At approximately 150 words per minute this script is about one minute; rehearse with the chosen demo clicks.
"""
    (OUT / "QUICK_PITCH.md").write_text(quick, encoding="utf-8")


if __name__ == "__main__":
    prepare_assets()
    prs, path = build_main()
    backup = build_backup()
    write_scripts()
    # Structural verification: every shape is on-canvas and editable text remains native.
    for s in prs.slides:
        for shape in s.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= prs.slide_width + 5
            assert shape.top + shape.height <= prs.slide_height + 5
    print(json.dumps({"pptx": str(path), "slides": len(prs.slides), "backup": str(backup), "format": "16:9", "editable_text_boxes": sum(sh.has_text_frame for s in prs.slides for sh in s.shapes)}))
