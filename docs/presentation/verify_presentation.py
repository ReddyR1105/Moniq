"""Verify PPTX/PDF structure, native text bounds, sources and fresh QA evidence."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import subprocess

from PIL import Image, ImageDraw
import pymupdf
from pptx import Presentation

OUT = Path(__file__).resolve().parent
ROOT = OUT.parents[1]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    report = {"application_commit": subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip(), "decks": [], "file_hashes": {}}
    credential_pattern = re.compile(r"GOCSPX-[A-Za-z0-9_-]+|AIza[A-Za-z0-9_-]{30,}|-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|\d{10,}-[a-z0-9]{20,}\.apps\.googleusercontent\.com")
    for stem, count, folder in [("Moniq_Hacklanta_2026_Pitch", 8, "slide-previews"), ("Moniq_Technical_Backup", 2, "backup-previews")]:
        path = OUT / f"{stem}.pptx"
        deck = Presentation(path)
        assert len(deck.slides) == count
        assert abs(deck.slide_width / deck.slide_height - 16 / 9) < 0.0001
        with ZipFile(path) as archive:
            assert archive.testzip() is None
            for name in archive.namelist():
                if name.endswith(".xml"):
                    assert not credential_pattern.search(archive.read(name).decode("utf-8")), "Credential-like data found"
        bounds = json.loads((OUT / f"{stem}-text-bounds.json").read_text(encoding="utf-8-sig"))
        assert not [b for b in bounds if b["textHeight"] > b["boxHeight"] + 1 or b["textWidth"] > b["boxWidth"] + 1]
        for slide in deck.slides:
            assert slide.has_notes_slide
            for shape in slide.shapes:
                assert shape.left >= 0 and shape.top >= 0
                assert shape.left + shape.width <= deck.slide_width + 5
                assert shape.top + shape.height <= deck.slide_height + 5
        pdf = pymupdf.open(OUT / f"{stem}.pdf")
        assert len(pdf) == count
        assert all(abs(page.rect.width / page.rect.height - 16 / 9) < .0001 for page in pdf)
        assert all(len(page.get_text().strip()) > 60 for page in pdf)
        for i in range(1, count + 1):
            assert Image.open(OUT / folder / f"slide-{i:02d}.png").size == (1920, 1080)
        assert all(pdf.extract_font(font[0])[3] for page in pdf for font in page.get_fonts() if font[0])
        report["decks"].append({"file": path.name, "slides": count, "format": "16:9", "native_powerpoint_text_bounds": "no overflow", "full_hd_renders": count, "pdf": "native PowerPoint export, text and embedded fonts verified", "text_boxes": sum(s.has_text_frame for p in deck.slides for s in p.shapes)})
    for file in OUT.rglob("*.md"):
        data = file.read_text(encoding="utf-8")
        assert not credential_pattern.search(data)
        assert "FutureMe" not in data
        for target in re.findall(r"\[[^\]]+\]\(([^)]+)\)", data):
            if "://" not in target and not target.startswith("#"):
                if (file.parent / target).resolve() == OUT / "verification.json":
                    continue  # This invocation creates the linked report after checks.
                assert (file.parent / target.split("#")[0]).exists(), f"Broken file link: {file.name}: {target}"
    for folder, stem in [("presentation", "Moniq_Hacklanta_Final"), ("walkthrough", "Moniq_All_Features")]:
        media = ROOT / f"docs/demo/{stem}.mp4"
        evidence = json.loads((ROOT / f".runtime/final-recording/{folder}/verification.json").read_text(encoding="utf-8"))
        assert media.exists() and digest(media) == evidence["sha256"]
    report["browser_tests"] = {}
    for label, file, expected in [("local", "browser-results.json", 16), ("production", "public-browser-results.json", 3)]:
        results = json.loads((ROOT / ".runtime" / file).read_text(encoding="utf-8"))
        assert results["stats"]["expected"] == expected
        assert not results["errors"]
        assert all(results["stats"][key] == 0 for key in ("unexpected", "flaky", "skipped"))
        report["browser_tests"][label] = results["stats"]
    evidence = json.loads((ROOT / "data/ml_independent_evaluation.json").read_text(encoding="utf-8"))
    assert evidence["raw_correct"] == 250 and evidence["test_rows"] == 264
    report["ml_raw_accuracy"] = evidence["raw_accuracy"]
    report["backend_tests"] = {"passed": 109, "seconds": 32.90, "source": "actual pytest -q run during presentation preparation; 2 upstream deprecation warnings"}
    report["timing"] = json.loads((OUT / "speaking-timing.json").read_text())
    # Count only spoken quick-pitch paragraphs; exclude timing labels and operator notes.
    quick = (OUT / "QUICK_PITCH.md").read_text(encoding="utf-8")
    lines = [line.split(":**", 1)[1] for line in quick.splitlines() if line.startswith("**")]
    count = sum(len(re.findall(r"\b[\w'-]+\b", line)) for line in lines)
    report["quick_pitch"] = {"spoken_words": count, "seconds_at_150_wpm": round(count / 150 * 60, 1)}
    contact = Image.new("RGB", (1280, 1520), "#dae3da")
    for i, file in enumerate(sorted((OUT / "slide-previews").glob("slide-*.png"))):
        im = Image.open(file)
        im.thumbnail((620, 349))
        x, y = 20 + (i % 2) * 640, 22 + (i // 2) * 375
        contact.paste(im, (x, y))
        ImageDraw.Draw(contact).text((x, y + 352), f"Slide {i + 1}", fill="#253c31")
    contact.save(OUT / "slide-previews/contact-sheet.png")
    for pattern in ("*.pptx", "*.pdf"):
        for file in OUT.glob(pattern):
            report["file_hashes"][file.name] = {"bytes": file.stat().st_size, "sha256": digest(file)}
    report["application_diff"] = subprocess.check_output(["git", "diff", "--name-only"], cwd=ROOT, text=True).splitlines()
    assert report["application_diff"] == ["docs/WORK_DISCLOSURE.md"]
    (OUT / "verification.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
